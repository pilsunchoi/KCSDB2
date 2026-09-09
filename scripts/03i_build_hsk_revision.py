"""
03i_build_hsk_revision.py — HSK 개정 이력 (dim_hsk_revision)

무엇을 푸는가:
  개정 목록은 어디에도 공표되지 않는다. 지금까지는 거래 자료에서 코드의 첫 등장·마지막
  등장이 몰리는 달을 세어 짐작했는데, 그 방법은 "코드가 없었다"와 "거래가 없었다"를
  가르지 못한다. 실제로 2012년 개정에서 신설된 968개 중 377개는 코드가 생긴 뒤 몇 해
  지나서야 거래가 붙었고, 90개는 끝내 거래가 없었다. 짐작으로는 이들을 놓친다.

두 경로를 쓴다:
  (1) 별표 차분 — 연도별 별표의 코드 집합을 빼면 신설·폐지가 정확히 나온다. 본줄기다.
  (2) 신구대조표 — 개정사유를 주고, 별표가 없는 구간을 갈라 준다. 검증에도 쓴다.

  둘이 서로를 검증한다. 2017년 별표와 2021년 별표의 차분은 신설 12·폐지 2인데,
  2019년과 2021년 신구대조표가 말하는 것을 더하면 신설 4+8·폐지 2+0으로 정확히 같다.

무엇을 주지 않는가:
  종전 코드가 어느 코드로 이어지는지는 여기 없다. 신구대조표에도 없다 - 좌우 양쪽에
  10단위 코드가 있으면서 서로 다른 행은 한 건도 없고, 변화는 언제나 삭제와 신설의
  짝이다. 대응관계는 여전히 추정(dim_hs10_to_2022)에 기댄다. 다만 개정사유가 있으면
  추정의 폭이 크게 줄어든다 - "무역량 감소에 따른 HSK 삭제"는 HS6가 그대로 남아
  물품이 같은 HS6 안의 남은 코드로 흡수되지만, 세계관세기구 개정은 HS6가 움직인다.

구간과 개정의 대응:
  별표 판본 사이에 개정이 둘 이상 낀 구간이 있다. 정직하게 적어 둔다.
    2011→2013 = 2012년 개정 + 2013년 개정(거래 자료 기준 신설 9건, 미미)
    2013→2015 = 2014년 개정 + 2015년 개정(신설 4건, 미미)
    2015→2017 = 2016년 개정(신설 1건) + 2017년 개정
    2017→2021 = 2018~2021년 개정. 신구대조표로 2019·2021을 갈라 담는다
    2021→2022 = 2022년 개정
    2022→2025 = 2025년 개정
  2019년 개정만 1월이 아니라 10월 시행이다. 신설된 마스크팩(3307904000)과
  탄소난방필름(8545902000)이 거래 자료에 둘 다 201910에 처음 나타난다.

담지 못한 것:
  2008·2009·2026년 개정. 별표는 2011년 것이 가장 이르고, 신구대조표는 재정경제부가
  보존기간 10년 경과로 폐기했으며 관세법령정보포털에도 없다. 관세법령정보포털
  세계HS > HS정보 > 관세율표에 한국 2002~2026년 판본이 매년 있으므로 거기서
  2007~2009년과 2026년 코드 집합을 얻으면 메울 수 있다.

입력:
  dim_hs10_name_hist (별표 2011·2013·2015·2017·2021)
  data/external/HSK_별표/HSK_별표_{2022,2025}.xlsx
  data/external/HSK_신구대조표/  (PDF·HWPX·XLSX)
출력: data/processed/kcsdb.duckdb 의 dim_hsk_revision

테이블 스키마:
  rev           VARCHAR   개정 이름. 개정 후 판본 연도
  effective_ym  INTEGER   시행 연월
  hs10          VARCHAR   10단위 품목번호
  change        VARCHAR   '신설' | '폐지'
  source        VARCHAR   '별표차분' | '신구대조표'
  name_ko       VARCHAR   품명(국문). 신설은 개정 후, 폐지는 개정 전 판본의 것
  reason        VARCHAR   개정사유. 2022·2025년만 있고 나머지는 NULL

실행:
  python scripts\\03i_build_hsk_revision.py
  python scripts\\03i_build_hsk_revision.py --dry-run   # DB 미기록, 추출·검증만
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import duckdb
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BYEOL_DIR = PROJECT_ROOT / "data" / "external" / "HSK_별표"
REV_DIR = PROJECT_ROOT / "data" / "external" / "HSK_신구대조표"
DB_PATH = PROJECT_ROOT / "data" / "processed" / "kcsdb.duckdb"
LOG_DIR = PROJECT_ROOT / "logs"
LOG_DIR.mkdir(exist_ok=True)

LOG_PATH = LOG_DIR / f"build_hsk_revision_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"), logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger(__name__)

NEW_TOKENS = ("신설", "신 설")
DEL_TOKENS = ("삭제", "삭 제", "폐지")

# 별표 차분으로 만드는 구간. (개정이름, 시행연월, 앞 판본, 뒤 판본)
SPANS = [
    ("2012", 201201, 2011, 2013),
    ("2014", 201401, 2013, 2015),
    ("2017", 201701, 2015, 2017),
    ("2022", 202201, 2021, 2022),
    ("2025", 202501, 2022, 2025),
]

# 2017→2021 구간은 별표만으로는 갈리지 않는다. 신구대조표가 둘로 나눈다.
SPLIT_SPAN = (2017, 2021)

# 신구대조표. 개정사유와 2019·2021 분리에 쓴다.
#   code_l/code_r : 코드 열의 (인덱스, 자릿수). 자릿수 합이 10이어야 10단위다.
#   carry         : 접두를 반복하지 않는 형식이면 True. 앞 행에서 물려받는다.
CONCORD = [
    dict(rev="2012", file="신구대조표_2011to2012.pdf", fmt="pdf",
         code_l=[(0, 4), (1, 2), (2, 2), (3, 2)], name_l=4,
         code_r=[(6, 4), (7, 2), (8, 2), (9, 2)], name_r=10,
         reason=None, carry=False, min_cols=11),
    dict(rev="2014", file="신구대조표_2013to2014.pdf", fmt="pdf",
         code_l=[(0, 4), (1, 2), (2, 4)], name_l=3,
         code_r=[(5, 4), (6, 2), (7, 4)], name_r=8,
         reason=None, carry=True, min_cols=10),
    dict(rev="2019", file="신구대조표_2017to2019.hwpx", fmt="hwpx",
         code_l=[(0, 4), (1, 2), (2, 2), (3, 2)], name_l=4,
         code_r=[(6, 4), (7, 2), (8, 2), (9, 2)], name_r=10,
         reason=None, carry=False, min_cols=11),
    dict(rev="2021", file="신구대조표_2019to2021.hwpx", fmt="hwpx",
         code_l=[(0, 4), (1, 2), (2, 2), (3, 2)], name_l=4,
         code_r=[(6, 4), (7, 2), (8, 2), (9, 2)], name_r=10,
         reason=None, carry=False, min_cols=11),
    dict(rev="2022", file="신구대조표_2021to2022.xlsx", fmt="xlsx",
         code_l=[(0, 4), (1, 2), (2, 2), (3, 2)], name_l=4,
         code_r=[(6, 4), (7, 2), (8, 2), (9, 2)], name_r=10,
         reason=12, carry=False, min_cols=13, header_rows=3),
    dict(rev="2025", file="신구대조표_2022to2025.xlsx", fmt="xlsx",
         code_l=[(0, 4), (1, 2), (2, 2), (3, 2)], name_l=4,
         code_r=[(6, 4), (7, 2), (8, 2), (9, 2)], name_r=10,
         reason=12, carry=False, min_cols=13, header_rows=4),
]
# 2015→2017 신구대조표(398쪽 PDF)는 괘선이 없어 표 인식이 되지 않는다. 30쪽을 훑어
# 열두 칸짜리 행이 여덟 개뿐이었다. 별표 차분으로 대신하고 사유는 포기한다.

EFFECTIVE = {"2012": 201201, "2014": 201401, "2017": 201701,
             "2019": 201910, "2021": 202101, "2022": 202201, "2025": 202501}


# ---------------------------------------------------------------- 별표 코드 집합

def byeolpyo_from_xlsx(path: Path) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(values_only=True):
        code = "".join((str(x).strip() if x is not None else "") for x in r[0:4])
        if len(code) == 10 and code.isdigit():
            rows.append((code, str(r[4]).strip() if len(r) > 4 and r[4] else None))
    wb.close()
    return pd.DataFrame(rows, columns=["hs10", "name_ko"]).drop_duplicates("hs10")


def load_byeolpyo(con) -> dict[int, pd.DataFrame]:
    out = {}
    for y in (2011, 2013, 2015, 2017, 2021):
        df = con.execute(
            "SELECT DISTINCT hs10, name_ko FROM dim_hs10_name_hist WHERE byeolpyo_year = ?", [y]
        ).df().drop_duplicates("hs10")
        out[y] = df
    for y in (2022, 2025):
        out[y] = byeolpyo_from_xlsx(BYEOL_DIR / f"HSK_별표_{y}.xlsx")
    return out


# ------------------------------------------------------------------- 신구대조표

def cell(row, i):
    if i is None or i >= len(row):
        return ""
    v = row[i]
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def read_pdf(path: Path):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in (page.extract_tables() or []):
                yield from table


def read_hwpx(path: Path):
    xml = ""
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if re.search(r"section\d+\.xml$", n):
                xml += z.read(n).decode("utf-8", "replace")
    for tr in re.findall(r"<hp:tr[ >].*?</hp:tr>", xml, re.S):
        cells = []
        for tc in re.findall(r"<hp:tc[ >].*?</hp:tc>", tr, re.S):
            t = "".join(re.findall(r"<hp:t>(.*?)</hp:t>", tc, re.S))
            t = re.sub(r"<[^>]+>", "", t)
            t = t.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
            cells.append(t.strip())
        yield cells


def read_xlsx(path: Path, header_rows: int):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= header_rows:
            yield list(row)
    wb.close()


def build_code(row, spec, carry):
    """코드 열을 이어 붙여 10단위 코드를 만든다.

    carry가 주어지면 빈 칸을 앞 행에서 물려받되, 위 단계가 바뀐 행에서는 아래 단계를
    지운다. 그래야 앞 항목의 하위 코드가 새 항목에 딸려오지 않는다.
    """
    parts = []
    changed = False
    for level, (idx, _width) in enumerate(spec):
        v = re.sub(r"\D", "", cell(row, idx))
        if carry is not None:
            if v:
                carry[level] = v
                changed = True
                for deeper in range(level + 1, len(spec)):
                    carry[deeper] = ""
            elif not changed:
                v = carry[level]
            else:
                v = ""
            if v:
                carry[level] = v
        parts.append(v)
    code = "".join(parts)
    return code if (len(code) == 10 and code.isdigit()) else ""


def extract_concord(src: dict) -> pd.DataFrame:
    path = REV_DIR / src["file"]
    if not path.exists():
        raise FileNotFoundError(path)

    if src["fmt"] == "pdf":
        rows = read_pdf(path)
    elif src["fmt"] == "hwpx":
        rows = read_hwpx(path)
    else:
        rows = read_xlsx(path, src.get("header_rows", 0))

    carry_l = [""] * len(src["code_l"]) if src["carry"] else None
    carry_r = [""] * len(src["code_r"]) if src["carry"] else None

    out = []
    for row in rows:
        if not row or len(row) < src["min_cols"]:
            continue
        lname, rname = cell(row, src["name_l"]), cell(row, src["name_r"])
        if lname == "품목번호" or rname == "품목번호":
            continue

        lcode = build_code(row, src["code_l"], carry_l)
        rcode = build_code(row, src["code_r"], carry_r)
        reason = cell(row, src["reason"]).split("\n")[0].strip() if src["reason"] else ""

        if lcode and (any(t in rname for t in DEL_TOKENS) or not rcode):
            out.append((src["rev"], "폐지", lcode, lname, reason))
        elif rcode and (any(t in lname for t in NEW_TOKENS) or not lcode):
            out.append((src["rev"], "신설", rcode, rname, reason))

    df = pd.DataFrame(out, columns=["rev", "change", "hs10", "name_ko", "reason"])
    return df.drop_duplicates(subset=["rev", "change", "hs10"])


# ------------------------------------------------------------------------ 본체

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="DB에 쓰지 않고 추출·검증만")
    args = ap.parse_args()

    con = duckdb.connect(str(DB_PATH), read_only=True)
    try:
        BP = load_byeolpyo(con)
    finally:
        con.close()
    for y, df in BP.items():
        logger.info("별표 %d: %d종", y, len(df))

    CD = pd.concat([extract_concord(s) for s in CONCORD], ignore_index=True)
    for rev, g in CD.groupby("rev"):
        logger.info("신구대조표 %s: 신설 %d 폐지 %d%s", rev,
                    (g["change"] == "신설").sum(), (g["change"] == "폐지").sum(),
                    " (사유 있음)" if g["reason"].str.len().gt(0).any() else "")

    # --- 별표 차분
    frames = []
    for rev, ym, y0, y1 in SPANS:
        s0, s1 = set(BP[y0]["hs10"]), set(BP[y1]["hs10"])
        new = BP[y1][BP[y1]["hs10"].isin(s1 - s0)].assign(change="신설")
        dele = BP[y0][BP[y0]["hs10"].isin(s0 - s1)].assign(change="폐지")
        d = pd.concat([new, dele], ignore_index=True)
        d["rev"], d["effective_ym"], d["source"] = rev, ym, "별표차분"
        frames.append(d)
        logger.info("별표차분 %s (%d→%d): 신설 %d 폐지 %d", rev, y0, y1, len(new), len(dele))

    # --- 2017→2021 구간은 신구대조표로 가른다
    y0, y1 = SPLIT_SPAN
    s0, s1 = set(BP[y0]["hs10"]), set(BP[y1]["hs10"])
    span_new, span_del = s1 - s0, s0 - s1
    sub = CD[CD["rev"].isin(["2019", "2021"])].copy()
    sub["effective_ym"] = sub["rev"].map(EFFECTIVE)
    sub["source"] = "신구대조표"
    frames.append(sub[["hs10", "name_ko", "change", "rev", "effective_ym", "source", "reason"]])

    got_new = set(sub.loc[sub["change"] == "신설", "hs10"])
    got_del = set(sub.loc[sub["change"] == "폐지", "hs10"])
    logger.info("2017→2021 별표차분 신설 %d 폐지 %d / 신구대조표(2019+2021) 신설 %d 폐지 %d",
                len(span_new), len(span_del), len(got_new), len(got_del))

    all_df = pd.concat(frames, ignore_index=True)
    if "reason" not in all_df:
        all_df["reason"] = None

    # --- 사유를 신구대조표에서 물려준다
    key = CD[CD["reason"].str.len() > 0][["rev", "change", "hs10", "reason"]]
    all_df = all_df.drop(columns=["reason"]).merge(key, on=["rev", "change", "hs10"], how="left")
    all_df = all_df[["rev", "effective_ym", "hs10", "change", "source", "name_ko", "reason"]]
    all_df = all_df.drop_duplicates(subset=["rev", "change", "hs10"]).reset_index(drop=True)

    # --- 검증
    assert all_df["hs10"].str.fullmatch(r"\d{10}").all(), "10자리가 아닌 코드가 있다"
    assert (all_df.groupby(["rev", "hs10"])["change"].nunique() == 1).all(), \
        "한 개정에서 신설이면서 폐지인 코드가 있다"
    assert got_new == span_new and got_del == span_del, \
        "2019·2021 신구대조표와 2017→2021 별표차분이 어긋난다"

    # 폐지 없이 신설만 있는 개정이 있다(2021년). 개정마다 행이 하나 이상이면 된다.
    per = all_df.groupby(["rev", "change"]).size().unstack(fill_value=0)
    assert set(per.index) == set(EFFECTIVE), f"개정이 빠졌다\n{per}"
    assert (per.sum(axis=1) > 0).all(), f"비어 있는 개정이 있다\n{per}"
    logger.info("\n%s", per.to_string())

    # --- 두 경로 일치율 (2019·2021 제외 — 그쪽은 신구대조표가 유일한 근거다)
    for rev in ("2012", "2014", "2022", "2025"):
        a = set(all_df.loc[(all_df["rev"] == rev), "hs10"])
        b = set(CD.loc[CD["rev"] == rev, "hs10"])
        if b:
            logger.info("%s 두 경로 겹침 %d / 별표차분 %d / 신구대조표 %d",
                        rev, len(a & b), len(a), len(b))

    # --- 거래 자료와 대조
    con = duckdb.connect(str(DB_PATH), read_only=True)
    try:
        con.register("rev_df", all_df)
        # 폐지 코드에 개정 뒤 거래가 잡히는 것은 대개 잔여 통관이다. 몇 달이나 이어지는지로
        # 가른다 - 석 달 안이면 통관 마무리, 한 해를 넘으면 같은 번호가 나중에 되살아난 것이다.
        chk = con.execute("""
            WITH f AS (SELECT hs10, MIN(yyyymm) fm, MAX(yyyymm) lm FROM fact_trade GROUP BY hs10),
            j AS (
                SELECT r.rev, r.change, r.effective_ym, f.fm, f.lm,
                       CASE WHEN r.change = '폐지' AND f.lm >= r.effective_ym
                            THEN (f.lm/100 - r.effective_ym/100)*12 + (f.lm%100 - r.effective_ym%100)
                       END AS 폐지후개월
                FROM rev_df r LEFT JOIN f USING (hs10)
            )
            SELECT rev, change, COUNT(*) AS 코드수, COUNT(fm) AS 거래있음,
                   SUM(CASE WHEN change='신설' AND fm <  effective_ym THEN 1 ELSE 0 END) AS 개정전거래,
                   SUM(CASE WHEN 폐지후개월 <= 3  THEN 1 ELSE 0 END) AS "폐지후 3개월내",
                   SUM(CASE WHEN 폐지후개월 > 12  THEN 1 ELSE 0 END) AS "폐지후 1년초과"
            FROM j GROUP BY 1,2 ORDER BY 1,2
        """).df()
        logger.info("\n거래 자료와 대조\n%s", chk.to_string(index=False))
    finally:
        con.close()

    if args.dry_run:
        logger.info("--dry-run: DB에 쓰지 않는다")
        return

    con = duckdb.connect(str(DB_PATH))
    try:
        con.register("rev_df", all_df)
        con.execute("DROP TABLE IF EXISTS dim_hsk_revision")
        con.execute("CREATE TABLE dim_hsk_revision AS SELECT * FROM rev_df")
        n = con.execute("SELECT COUNT(*) FROM dim_hsk_revision").fetchone()[0]
        logger.info("dim_hsk_revision %d행 적재", n)
    finally:
        con.close()


if __name__ == "__main__":
    main()
